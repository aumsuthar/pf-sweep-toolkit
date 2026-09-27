"""Build CM_sweep.xlsx: one sheet per sweep CSV plus a combined sheet,
with the ParaView CM render embedded in each design point's row.

Sources: <stem>.csv (sweep results) + <stem>_CM/manifest.csv (CM parameters
and render paths, written by make_cm.py).

Usage:  python3 make_cm_excel.py set1 set2
"""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.chart.marker import Marker
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageChops

BASE      = Path(__file__).resolve().parent
THUMB_PX  = 130                      # thumbnail height in pixels
ROW_PT    = THUMB_PX * 0.75 + 6      # Excel row height (points) with a little padding
IMG_COL_W = 21                       # Excel column width units

HEAD_FILL = PatternFill("solid", fgColor="2A1A0E")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
THIN      = Side(style="thin", color="D9D2C8")
BORDER    = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# (header, key, number format, column width)
COLUMNS = [
    ("Set",             "set",             "@",       7),
    ("Point",           "point",           "0",       7),
    ("CM render",       "_image",          "@",       IMG_COL_W),
    ("avg ⟨c⟩ (%)",     "avg",             "0.00",    11),
    ("λ (nm)",          "wv",              "0",       9),
    ("waves / 640 nm",  "waves_in_domain", "0.00",    14),
    ("amp (±%)",        "amp",             "0.00",    10),
    ("c min (%)",       "c_min_pct",       "0.000",   10),
    ("c max (%)",       "c_max_pct",       "0.000",   10),
    # --- the four ML metrics, each paired with its phase-field analogue ---
    ("ML V1 stress",       "V1_stress",      "0",      12),
    ("SIM V1 stress",      "sim_V1",         "0",      12),
    ("\u0394 V1",              "d_V1",           "+0;-0",  9),
    ("ML V3 stress",       "V3_stress",      "0",      12),
    ("SIM V3 stress",      "sim_V3",         "0",      12),
    ("\u0394 V3",              "d_V3",           "+0;-0",  9),
    ("ML strain range",    "strain_range",   "0.00",   14),
    ("SIM strain range",   "sim_strain_rng", "0.00",   14),
    ("\u0394 strain (pp)",     "d_strain_pp",    "+0.00;-0.00", 14),
    ("ML loop area",       "loop_area",      "0.0",    12),
    ("SIM loop area",      "sim_loop_area",  "0.00",   12),
    ("ML/SIM loop ratio",  "loop_ratio",     "0.0",    16),
    ("loop rank \u0394",       "d_loop_rank",    "+0;-0",  11),
    # --- supporting detail ---
    ("SIM max vf",         "sim_max_vf",     "0.000",  11),
    ("SIM cliff \u0394vf",      "sim_max_vf_drop","0.000",  13),
    ("SIM residual strain","sim_residual_strain","0.000000",15),
    ("SIM notes",          "sim_note",       "@",      24),
    ("CM field (.vtk)", "_vtk",            "@",       36),
]


def crop_cube(path: Path) -> Image.Image:
    """Drop the colour bar and header text from a render, trim to the cube."""
    im = Image.open(path).convert("RGB")
    w, h = im.size
    im = im.crop((0, int(h * 0.09), int(w * 0.88), h))
    bbox = ImageChops.difference(im, Image.new("RGB", im.size, (255, 255, 255))).getbbox()
    if bbox:
        im = im.crop(bbox)
    scale = THUMB_PX / im.height
    return im.resize((max(1, round(im.width * scale)), THUMB_PX), Image.LANCZOS)


import os
SIM_DIR = os.environ.get("SIM_DIR", "set1_simulation_results")
OUT_NAME = os.environ.get("OUT_NAME", "CM_sweep.xlsx")
FIG_SUFFIX = os.environ.get("FIG_SUFFIX", "")
SIM_CSV = BASE.parent / SIM_DIR / "set1_summary.csv"
SIM_FIELDS = ["sim_forward_MPa", "sim_reverse_MPa", "sim_top_MPa", "sim_max_strain",
              "sim_max_vf", "sim_residual_strain", "sim_note", "sim_loop_area",
              "sim_V1", "sim_V3", "sim_strain_rng", "sim_strain_rng_at1100", "sim_max_vf_drop"]


def load_sim(stem: str) -> dict[int, dict]:
    """Phase-field results, keyed by point. Only set1 has been simulated."""
    if stem != "set1" or not SIM_CSV.exists():
        return {}
    out = {}
    for r in csv.DictReader(open(SIM_CSV)):
        out[int(r["point"])] = {
            k: (float(r[k]) if k in ("sim_max_strain", "sim_max_vf", "sim_residual_strain",
                                     "sim_loop_area", "sim_V1", "sim_V3", "sim_strain_rng",
                                     "sim_strain_rng_at1100", "sim_max_vf_drop")
                else int(r[k]) if k == "sim_top_MPa"
                else r[k])
            for k in SIM_FIELDS if k in r
        }
    return out


def load_set(stem: str) -> list[dict]:
    """Merge sweep CSV rows with their CM manifest entry, newest thumbnail cached."""
    cm_dir = BASE / f"{stem}_CM"
    sweep  = {int(float(r["point"])): r for r in csv.DictReader(open(BASE / f"{stem}.csv"))}
    man    = {int(float(r["point"])): r for r in csv.DictReader(open(cm_dir / "manifest.csv"))}
    thumbs = cm_dir / "thumbs"
    thumbs.mkdir(exist_ok=True)

    sim = load_sim(stem)

    blank = {k: ("" if k in ("sim_forward_MPa", "sim_reverse_MPa", "sim_note") else None)
             for k in SIM_FIELDS}

    rows = []
    for pt in sorted(sweep):
        s, m = sweep[pt], man[pt]
        thumb = thumbs / f"point{pt:02d}.png"
        crop_cube(cm_dir / m["png"]).save(thumb)
        rows.append({
            "set": stem,
            "point": pt,
            "avg": float(s["avg"]),
            "wv": float(s["wv"]),
            "amp": float(s["amp"]),
            "waves_in_domain": float(m["waves_in_domain"]),
            "c_min_pct": float(m["c_min_pct"]),
            "c_max_pct": float(m["c_max_pct"]),
            "V1_stress": float(s["V1_stress"]),
            "V3_stress": float(s["V3_stress"]),
            "strain_range": float(s["strain_range"]),
            "loop_area": float(s["loop_area"]),
            **sim.get(pt, blank),
            "_thumb": thumb,
            "_vtk": f"{stem}_CM/{m['vtk']}",
        })
    add_comparison(rows)
    return rows


def _mid(window: str):
    """Midpoint of a "750-1000" transformation window."""
    try:
        lo, hi = window.split("-")
        return (float(lo) + float(hi)) / 2.0
    except Exception:
        return None


def add_comparison(rows: list[dict]) -> None:
    """Pair each ML metric with its phase-field analogue.

    Mapping used for V1/V3 (the ML's own definition is not documented): in both
    data sets V1 < V3 for every point, and the phase field gives a forward
    (upper) and reverse (lower) transformation stress, so
        V1  <->  reverse / lower plateau
        V3  <->  forward / upper plateau
    Treat the V1/V3 comparison as provisional until that mapping is confirmed.

    Loop area is compared as a RATIO, not a difference: the two differ by ~74x,
    so subtracting them would only restate the units gap. A constant ratio would
    mean the ML is right up to a conversion factor; the spread is the real
    disagreement.
    """
    blanks = dict(d_V1=None, d_V3=None, d_strain_pp=None,
                  loop_ratio=None, d_loop_rank=None)
    have = [r for r in rows if r.get("sim_loop_area") not in (None, "")]
    if not have:
        for r in rows:
            r.update(blanks)
        return

    ml_rank = {r["point"]: i + 1 for i, r in
               enumerate(sorted(have, key=lambda r: r["loop_area"]))}
    sim_rank = {r["point"]: i + 1 for i, r in
                enumerate(sorted(have, key=lambda r: r["sim_loop_area"]))}

    for r in rows:
        if r.get("sim_loop_area") in (None, ""):
            r.update(blanks)
            continue
        r["d_V1"] = round(r["sim_V1"] - r["V1_stress"], 1)
        r["d_V3"] = round(r["sim_V3"] - r["V3_stress"], 1)
        r["d_strain_pp"] = round(r["sim_strain_rng"] - r["strain_range"], 4)
        r["loop_ratio"] = round(r["loop_area"] / r["sim_loop_area"], 2)
        r["d_loop_rank"] = ml_rank[r["point"]] - sim_rank[r["point"]]


def write_sheet(ws, rows, show_set_col: bool):
    cols = [c for c in COLUMNS if show_set_col or c[1] != "set"]

    for i, (header, _key, _fmt, width) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i, value=header)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 30

    for r, row in enumerate(rows, start=2):
        ws.row_dimensions[r].height = ROW_PT
        for i, (_h, key, fmt, _w) in enumerate(cols, start=1):
            cell = ws.cell(row=r, column=i)
            cell.border = BORDER
            if key == "_image":
                img = XLImage(row["_thumb"])
                img.anchor = f"{get_column_letter(i)}{r}"
                ws.add_image(img)
                continue
            val = row[key]
            cell.value = "" if val is None else val
            cell.number_format = fmt
            cell.alignment = Alignment(
                horizontal="left" if key in ("set", "_vtk") else "center",
                vertical="center")
            if key == "_vtk":
                cell.font = Font(size=8, color="7A5548")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"



# Categorical slots 1 and 2 of the validated reference palette. Checked with
# the data-viz validator: CVD dE 24.7 (protan), normal-vision dE 33.6, both
# pass contrast vs surface. ML is always slot 1, SIM always slot 2 — colour
# follows the entity, never the rank.
C_ML, C_SIM, C_NEUTRAL = "2A78D6", "EB6834", "7A7A75"


def _style(chart, title, x_title, y_title, legend=True):
    chart.title = title
    chart.x_axis.title = x_title
    chart.y_axis.title = y_title
    chart.height, chart.width = 8.5, 14
    chart.x_axis.majorGridlines = None           # recessive axes: no vertical grid
    chart.y_axis.delete = False
    chart.x_axis.delete = False
    if legend:
        chart.legend.position = "t"
        chart.legend.overlay = False
    else:
        chart.legend = None                      # one series: the title names it


def _line(series, hex_colour, width=20000):
    series.graphicalProperties.line.solidFill = hex_colour
    series.graphicalProperties.line.width = width
    series.marker = Marker(symbol="circle", size=8)
    series.marker.graphicalProperties.solidFill = hex_colour
    series.marker.graphicalProperties.line.solidFill = hex_colour


def add_charts(wb, rows):
    """A Charts tab: the matplotlib figures plus the numbers behind them.

    These are rendered PNGs rather than native Excel charts. openpyxl's charts
    are sized in cell units and its axis labels collide at this data density —
    the rendered figures control figure size, marker size and label placement,
    so they stay readable. Regenerate with ../plot_set1.py, then rerun this.
    """
    have = [r for r in rows if r.get("sim_loop_area") not in (None, "")]
    if not have:
        return
    ws = wb.create_sheet("Charts")
    ws.sheet_view.showGridLines = False

    figs = [(f"set1_hysteresis_loops{FIG_SUFFIX}.png", "Simulated superelastic cycles", 1500),
            (f"set1_ml_vs_sim{FIG_SUFFIX}.png",        "ML prediction vs phase field",  1080)]
    row = 1
    for fname, caption, width_px in figs:
        path = BASE.parent / fname
        if not path.exists():
            continue
        c = ws.cell(row=row, column=1, value=caption)
        c.font = Font(bold=True, size=12, color="2A1A0E")
        img = XLImage(str(path))
        scale = width_px / img.width
        img.width, img.height = width_px, int(img.height * scale)
        img.anchor = f"A{row + 1}"
        ws.add_image(img)
        row += int(img.height / 19) + 4        # ~19 px per default row

    # the numbers behind the figures, so the tab is self-contained
    c = ws.cell(row=row, column=1, value="Underlying values")
    c.font = Font(bold=True, size=12, color="2A1A0E")
    row += 1
    hdr = ["Point", "ML V1", "SIM V1", "ML V3", "SIM V3",
           "ML strain range", "SIM strain range", "ML/SIM loop ratio"]
    for i, h in enumerate(hdr, start=1):
        cell = ws.cell(row=row, column=i, value=h)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = 17
    for j, r in enumerate(sorted(have, key=lambda r: r["point"]), start=row + 1):
        for i, (v, fmt) in enumerate([(r["point"], "0"), (r["V1_stress"], "0"),
                                      (r["sim_V1"], "0"), (r["V3_stress"], "0"),
                                      (r["sim_V3"], "0"), (r["strain_range"], "0.00"),
                                      (r["sim_strain_rng"], "0.00"),
                                      (r["loop_ratio"], "0.0")], start=1):
            cell = ws.cell(row=j, column=i, value=v)
            cell.number_format = fmt
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center")



def add_raw_data(wb):
    """Every simulated step, one row: the data behind every figure and metric."""
    src = BASE.parent / SIM_DIR / "set1_results.csv"
    if not src.exists():
        return
    rows = list(csv.DictReader(open(src)))
    # CM, then loading before unloading, then ascending stress
    rows.sort(key=lambda r: (r["cm"], 0 if r["branch"].startswith("load") else 1,
                             float(r["stress_mpa"])))
    ws = wb.create_sheet("Raw data")

    cols = [("CM", "cm", "@", 8), ("\u03bb (nm)", "lambda_nm", "0", 9),
            ("amp (\u00b1%)", "amp", "0.00", 10), ("avg \u27e8c\u27e9 (%)", "avg", "0.00", 11),
            ("branch", "branch", "@", 10), ("step label", "label", "@", 14),
            ("stress (MPa)", "stress_mpa", "0.0", 13),
            ("strain", "strain", "0.000000", 12),
            ("strain (%)", None, "0.0000", 11),
            ("vf", "vf", "0.000000", 11), ("v1", "v1", "0.000000", 11),
            ("v2", "v2", "0.000000", 11), ("v3", "v3", "0.000000", 11),
            ("v4", "v4", "0.000000", 11)]

    for i, (h, _k, _f, w) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 26

    num = {"lambda_nm", "amp", "avg", "stress_mpa", "strain", "vf", "v1", "v2", "v3", "v4"}
    for r, row in enumerate(rows, start=2):
        for i, (_h, k, fmt, _w) in enumerate(cols, start=1):
            cell = ws.cell(row=r, column=i)
            if k is None:                                  # derived percent column
                cell.value = float(row["strain"]) * 100.0
            else:
                cell.value = float(row[k]) if k in num else row[k]
            cell.number_format = fmt
            cell.alignment = Alignment(horizontal="center" if k != "label" else "left")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
    print(f"  raw data: {len(rows)} rows")


def main(stems):
    wb = Workbook()
    wb.remove(wb.active)
    combined = []
    for stem in stems:
        rows = load_set(stem)
        combined += rows
        write_sheet(wb.create_sheet(stem), rows, show_set_col=False)
        if stem == "set1":
            set1_rows = rows
        print(f"  {stem}: {len(rows)} rows")
    write_sheet(wb.create_sheet("combined"), combined, show_set_col=True)
    try:
        add_raw_data(wb)
    except Exception as e:
        print("  raw data skipped:", e)
    try:
        add_charts(wb, set1_rows)
        print("  charts: added")
    except NameError:
        pass
    print(f"  combined: {len(combined)} rows")

    out = BASE / OUT_NAME
    wb.save(out)
    print("saved", out)


if __name__ == "__main__":
    main(sys.argv[1:] or ["set1", "set2"])
